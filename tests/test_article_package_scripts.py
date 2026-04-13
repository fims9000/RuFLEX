from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def test_article_readiness_script_generates_reports() -> None:
    build_result = subprocess.run(
        [sys.executable, "scripts/build_article_final_package.py", "--skip-suite"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    assert build_result.returncode == 0, build_result.stderr or build_result.stdout

    result = subprocess.run(
        [sys.executable, "scripts/check_article_readiness.py"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr or result.stdout

    report_path = Path(result.stdout.strip())
    assert report_path.exists()

    json_path = ROOT / "docs/article/readiness_report.json"
    assert json_path.exists()

    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert payload["overall_status"] == "pass"
    assert payload["submission_status"] == "pending"
    assert payload["profile_status"] == "template_placeholders_detected"
    assert payload["references_status"] == "ready"
    assert any(item["key"] == "profile_confirmed" for item in payload["submission_items"])
    assert payload["profile_placeholder_warnings"]
    assert any(item["name"] == "submission_bundle_zip" and item["ok"] for item in payload["checks"])
    assert any(item["name"] == "article_profile_valid" and item["ok"] for item in payload["checks"])
    assert any(item["name"] == "article_references_valid" and item["ok"] for item in payload["checks"])
    assert any(item["name"] == "submission_state_valid" and item["ok"] for item in payload["checks"])
    assert any(item["name"] == "doc:docs/article/shablon_dokladov_ready.docx" and item["ok"] for item in payload["checks"])
    assert any(
        item["name"] == "doc:docs/article/conference_template_ready_en.docx" and item["ok"]
        for item in payload["checks"]
    )
    assert any(
        item["name"] == "doc:docs/article/shablon_dokladov_illustrated.docx" and item["ok"]
        for item in payload["checks"]
    )
    assert any(
        item["name"] == "doc:docs/article/conference_template_illustrated_en.docx" and item["ok"]
        for item in payload["checks"]
    )

    bundle_report = ROOT / "docs/article/submission_bundle/docs/readiness_report.md"
    assert bundle_report.exists()
    assert (ROOT / "docs/article/submission_bundle/docs/article_references.json").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/submission_state.json").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/submission_state.template.json").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/references_ru_gost.md").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/references_en_ieee.md").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/shablon_dokladov_ready.docx").exists()
    assert (ROOT / "docs/article/submission_bundle/docs/conference_template_ready_en.docx").exists()
    ru_illustrated = ROOT / "docs/article/submission_bundle/docs/shablon_dokladov_illustrated.docx"
    en_illustrated = ROOT / "docs/article/submission_bundle/docs/conference_template_illustrated_en.docx"
    ru_illustrated_pdf = ROOT / "docs/article/submission_bundle/docs/shablon_dokladov_illustrated.pdf"
    en_illustrated_pdf = ROOT / "docs/article/submission_bundle/docs/conference_template_illustrated_en.pdf"
    paper_pdf = ROOT / "docs/article/submission_bundle/docs/paper_draft.pdf"
    rinc_pdf = ROOT / "docs/article/submission_bundle/docs/rinc_draft.pdf"
    assert ru_illustrated.exists()
    assert en_illustrated.exists()
    assert ru_illustrated_pdf.exists()
    assert en_illustrated_pdf.exists()
    assert paper_pdf.exists()
    assert rinc_pdf.exists()

    with zipfile.ZipFile(ru_illustrated) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())
    with zipfile.ZipFile(en_illustrated) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())
    assert ru_illustrated_pdf.read_bytes().startswith(b"%PDF")
    assert en_illustrated_pdf.read_bytes().startswith(b"%PDF")
    assert paper_pdf.read_bytes().startswith(b"%PDF")
    assert rinc_pdf.read_bytes().startswith(b"%PDF")


def test_illustrated_manuscripts_pick_up_article_profile() -> None:
    profile_path = ROOT / "docs/article/article_profile.json"
    original_text = profile_path.read_text(encoding="utf-8")
    payload = json.loads(original_text)
    payload["authors"][0]["name_ru"] = "Тестовый Автор"
    payload["authors"][0]["name_en"] = "Test Author"
    payload["authors"][0]["email"] = "test.author@example.com"
    payload["authors"][0]["orcid"] = "1111-2222-3333-4444"
    payload["funding_ru"] = "Тестовое финансирование."
    payload["funding_en"] = "Test funding."

    try:
        profile_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, "scripts/build_article_illustrated_manuscripts.py"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        assert result.returncode == 0, result.stderr or result.stdout

        ru_doc = Document(ROOT / "docs/article/shablon_dokladov_illustrated.docx")
        en_doc = Document(ROOT / "docs/article/conference_template_illustrated_en.docx")
        ru_text = "\n".join(paragraph.text for paragraph in ru_doc.paragraphs)
        en_text = "\n".join(paragraph.text for paragraph in en_doc.paragraphs)
        profile_card = (ROOT / "docs/article/article_profile_card.md").read_text(encoding="utf-8")
        final_metadata = (ROOT / "docs/article/final_metadata.md").read_text(encoding="utf-8")
        ready_ru = (ROOT / "docs/article/shablon_dokladov_ready.md").read_text(encoding="utf-8")
        ready_en = (ROOT / "docs/article/conference_template_ready_en.md").read_text(encoding="utf-8")
        references_ru = (ROOT / "docs/article/references_ru_gost.md").read_text(encoding="utf-8")
        references_en = (ROOT / "docs/article/references_en_ieee.md").read_text(encoding="utf-8")
        rinc_md = (ROOT / "docs/article/rinc_draft.md").read_text(encoding="utf-8")
        rinc_txt = (ROOT / "docs/article/rinc_draft.txt").read_text(encoding="utf-8")

        assert "Тестовый Автор" in ru_text
        assert "test.author@example.com" in ru_text
        assert "1111-2222-3333-4444" in ru_text
        assert "Тестовое финансирование." in ru_text
        assert "Test Author" in en_text
        assert "test.author@example.com" in en_text
        assert "1111-2222-3333-4444" in en_text
        assert "Test funding." in en_text
        assert "Тестовый Автор" in profile_card
        assert "test.author@example.com" in profile_card
        assert "1111-2222-3333-4444" in profile_card
        assert "Тестовый Автор" in ready_ru
        assert "Test Author" in ready_en
        assert "test.author@example.com" in ready_ru
        assert "test.author@example.com" in ready_en
        assert "1111-2222-3333-4444" in ready_ru
        assert "1111-2222-3333-4444" in ready_en
        assert "1111-2222-3333-4444" in rinc_md
        assert "1111-2222-3333-4444" in rinc_txt
        assert payload["title_ru"] in final_metadata
        assert payload["title_en"] in final_metadata
        assert "ANFIS" in references_ru
        assert "PyTorch" in references_en
        assert "Тестовый Автор" in rinc_md
        assert "Test Author" in rinc_txt
    finally:
        profile_path.write_text(original_text, encoding="utf-8")
        subprocess.run(
            [sys.executable, "scripts/render_article_profile_docs.py"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        subprocess.run(
            [sys.executable, "scripts/build_article_illustrated_manuscripts.py"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )


def test_strict_submission_gate_fails_for_placeholder_profile() -> None:
    result = subprocess.run(
        [sys.executable, "scripts/check_article_readiness.py", "--strict-submission"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode != 0
    assert "strict-submission gate failed" in (result.stderr or result.stdout)


def test_strict_submission_gate_passes_for_realistic_profile_and_completed_state() -> None:
    profile_path = ROOT / "docs/article/article_profile.json"
    state_path = ROOT / "docs/article/submission_state.json"
    original_profile = profile_path.read_text(encoding="utf-8")
    original_state = state_path.read_text(encoding="utf-8")
    payload = json.loads(original_profile)
    state_payload = json.loads(original_state)

    payload["authors"][0].update(
        {
            "name_ru": "Иван Иванов",
            "name_en": "Ivan Ivanov",
            "email": "ivan.ivanov@example.org",
            "orcid": "0000-0001-2345-6789",
            "affiliation_ru": "Университет, Москва, Россия",
            "affiliation_en": "University, Moscow, Russia",
        }
    )
    payload["authors"][1].update(
        {
            "name_ru": "Петр Петров",
            "name_en": "Petr Petrov",
            "email": "petr.petrov@example.org",
            "orcid": "0000-0002-3456-7890",
            "affiliation_ru": "Институт, Санкт-Петербург, Россия",
            "affiliation_en": "Institute, Saint Petersburg, Russia",
        }
    )
    state_payload.update(
        {
            "profile_confirmed": True,
            "references_checked": True,
            "figures_finalized": True,
            "antiplagiat_screenshot_added": True,
            "layout_reviewed": True,
            "notes": "Completed for strict submission gate test.",
        }
    )

    try:
        profile_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        state_path.write_text(json.dumps(state_payload, indent=2, ensure_ascii=False), encoding="utf-8")

        build_result = subprocess.run(
            [sys.executable, "scripts/build_article_final_package.py", "--skip-suite"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        assert build_result.returncode == 0, build_result.stderr or build_result.stdout

        result = subprocess.run(
            [sys.executable, "scripts/check_article_readiness.py", "--strict-submission"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
        assert result.returncode == 0, result.stderr or result.stdout

        report_payload = json.loads((ROOT / "docs/article/readiness_report.json").read_text(encoding="utf-8"))
        assert report_payload["profile_status"] == "ready"
        assert report_payload["submission_status"] == "ready"
    finally:
        profile_path.write_text(original_profile, encoding="utf-8")
        state_path.write_text(original_state, encoding="utf-8")
        subprocess.run(
            [sys.executable, "scripts/build_article_final_package.py", "--skip-suite"],
            cwd=ROOT,
            capture_output=True,
            text=True,
            check=False,
        )
