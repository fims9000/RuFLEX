from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from article_profile import (
    format_affiliations,
    format_author_names,
    format_emails,
    format_keywords,
    load_article_profile,
)
from render_article_profile_docs import render_article_profile_docs

ROOT = Path(__file__).resolve().parents[1]
FIGURES_DIR = ROOT / "docs/article/submission_bundle/figures"
TABLES_DIR = ROOT / "docs/article/submission_bundle/tables"

RU_OUTPUT = ROOT / "docs/article/shablon_dokladov_illustrated.docx"
EN_OUTPUT = ROOT / "docs/article/conference_template_illustrated_en.docx"


def main() -> None:
    build_illustrated_manuscripts()
    print(RU_OUTPUT)
    print(EN_OUTPUT)


def build_illustrated_manuscripts() -> tuple[Path, Path]:
    render_article_profile_docs()
    profile = load_article_profile()
    regression_rows = _read_csv_rows(TABLES_DIR / "regression_results_table.csv")
    classification_rows = _read_csv_rows(TABLES_DIR / "classification_results_table.csv")

    _build_russian_manuscript(
        profile=profile,
        regression_rows=regression_rows,
        classification_rows=classification_rows,
    )
    _build_english_manuscript(
        profile=profile,
        regression_rows=regression_rows,
        classification_rows=classification_rows,
    )
    return RU_OUTPUT, EN_OUTPUT


def _build_russian_manuscript(
    *,
    profile: dict[str, object],
    regression_rows: list[dict[str, str]],
    classification_rows: list[dict[str, str]],
) -> None:
    document = Document()
    _set_default_font(document, "Times New Roman", 11)

    _add_centered_title(
        document,
        str(profile["title_ru"]),
    )
    _add_centered_text(
        document,
        [
            format_author_names(profile, language="ru"),
            *format_affiliations(profile, language="ru"),
            f"E-mail: {format_emails(profile)}",
        ],
    )

    _add_heading(document, "Аннотация")
    _add_paragraph(
        document,
        "В работе представлена платформа RuFLEX для гибридного глубокого нечеткого обучения, "
        "ориентированная на задачи геоаналитики, пространственно обусловленной классификации, "
        "прогнозирования и интеллектуальной поддержки принятия решений. В отличие от black-box-подходов, "
        "платформа объединяет интерпретируемые rule-based механизмы, дифференцируемые нечеткие слои и "
        "архитектуры deep fuzzy feature learning в едином контуре, реализованном на Python и PyTorch. "
        "Экспериментальная валидация проведена на двух geo-oriented табличных задачах: California Housing Regression "
        "и California Value Binary Geo. На regression-задаче лучший результат показала конфигурация Deep Research, "
        "а на binary classification лучшей оказалась конфигурация Flat Interpretable."
    )
    _add_paragraph(
        document,
        f"Ключевые слова: {format_keywords(profile['keywords_ru'])}",
        italic_prefix="Ключевые слова:",
    )

    _add_heading(document, "1. Введение")
    _add_paragraph(
        document,
        "RuFLEX позиционируется как платформенный слой для гибридного глубокого нечеткого обучения, а не как одна "
        "фиксированная модель. Платформа развивает собственные SDK/UI/explainability-компоненты поверх vendored "
        "deep fuzzy backend `ruanfis`, происходящего из репозитория `deep-neuro-fuzzy`. Такое разделение позволяет "
        "сохранять преемственность deep fuzzy ядра и одновременно развивать собственный исследовательский и "
        "прикладной контур."
    )
    _add_paragraph(
        document,
        "Ключевая идея работы состоит в объединении плоских интерпретируемых neuro-fuzzy baseline-конфигураций и "
        "режима deep fuzzy feature learning в единой воспроизводимой платформе, которая поддерживает не только "
        "обучение и инференс, но и анализ функций принадлежности, правил, скрытых концептов и вклада факторов."
    )

    _add_heading(document, "2. Архитектура платформы и модельные режимы")
    _add_paragraph(
        document,
        "Текущая архитектура RuFLEX включает ядро платформы, модельный слой, слой обучения, explainability layer и "
        "инструментальный слой. В первой версии реализованы два базовых режима: flat neuro-fuzzy baseline и deep "
        "fuzzy feature learning. Первый режим удобен для компактного и интерпретируемого предсказания, второй "
        "использует скрытые слои как механизм формирования новых нечетких концептов."
    )
    _add_bullets(
        document,
        [
            "Python SDK и toolbox-style API",
            "visual workbench на базе Streamlit",
            "stage-wise pretraining, refinement и fine-tuning",
            "Explainability Dashboard с анализом membership functions, rules и hidden concepts",
            "reproducible article benchmark pipeline",
        ],
    )

    _add_heading(document, "3. Экспериментальная постановка")
    _add_paragraph(
        document,
        "Для валидации использованы две geo-oriented табличные задачи. Первая задача, California Housing Regression, "
        "основана на `sklearn.fetch_california_housing` и использует пространственные признаки `Latitude` и "
        "`Longitude`. Вторая задача, California Value Binary Geo, построена на том же наборе данных, но в постановке "
        "сбалансированной бинарной классификации по признаку `HighValue`."
    )
    _add_paragraph(
        document,
        "Во всех экспериментах сравнивались конфигурации Flat Baseline, Flat Interpretable, Deep Article Demo и "
        "Deep Research. Все результаты были получены в воспроизводимом article suite-контуре."
    )

    _add_heading(document, "4. Результаты")
    _add_paragraph(
        document,
        "На задаче California Housing Regression лучшей конфигурацией оказалась Deep Research "
        "(`test_rmse = 0.795757`, `test_r2 = 0.507159`). На задаче California Value Binary Geo лучшей конфигурацией "
        "оказалась Flat Interpretable (`test_accuracy = 0.746341`, `test_f1 = 0.739348`). Тем самым разные режимы "
        "RuFLEX оказываются предпочтительными на разных типах задач, что поддерживает именно платформенное "
        "позиционирование системы."
    )

    _add_metric_table(
        document,
        title="Таблица 1. Сравнение конфигураций на regression-задаче",
        headers=("Конфигурация", "Test RMSE", "Test MAE", "Test R2"),
        rows=[
            (
                row["variant_label"],
                _fmt(row["test_rmse"]),
                _fmt(row["test_mae"]),
                _fmt(row["test_r2"]),
            )
            for row in regression_rows
        ],
    )
    _add_metric_table(
        document,
        title="Таблица 2. Сравнение конфигураций на classification-задаче",
        headers=("Конфигурация", "Test Accuracy", "Test Precision", "Test F1"),
        rows=[
            (
                row["variant_label"],
                _fmt(row["test_accuracy"]),
                _fmt(row["test_precision"]),
                _fmt(row["test_f1"]),
            )
            for row in classification_rows
        ],
    )

    _add_heading(document, "5. Иллюстративные материалы explainability")
    _add_figure(
        document,
        FIGURES_DIR / "figure1_regression_benchmark.png",
        "Рис. 1. Сравнение flat и deep fuzzy конфигураций на задаче California Housing Regression.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure2_regression_board.png",
        "Рис. 2. Explainability board для лучшей regression-конфигурации.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure3_classification_benchmark.png",
        "Рис. 3. Сравнение конфигураций RuFLEX на задаче California Value Binary Geo.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure4_classification_board.png",
        "Рис. 4. Explainability board для лучшей classification-конфигурации.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure5_membership_example.png",
        "Рис. 5. Пример функций принадлежности для интерпретируемой фаззификации входного признака.",
        width=5.7,
    )

    _add_heading(document, "6. Заключение")
    _add_paragraph(
        document,
        "Разработана платформа RuFLEX для гибридного глубокого нечеткого обучения, объединяющая интерпретируемые "
        "rule-based механизмы, дифференцируемые нечеткие слои, deep fuzzy feature learning, SDK, visual workbench "
        "и explainability tooling в едином reproducible контуре. Показано, что платформа работает как на "
        "regression-, так и на classification-сценариях, а выбор лучшей конфигурации зависит от структуры задачи."
    )
    _add_paragraph(
        document,
        "Перспективы дальнейшего развития включают mature spatial/deep fuzzy module, более развитую logical "
        "traceability engine, counterfactual explainability и перенос вычислительно тяжелых частей в C++."
    )

    _add_heading(document, "Благодарность")
    _add_paragraph(document, str(profile["acknowledgment_ru"]))
    _add_heading(document, "Финансирование")
    _add_paragraph(document, str(profile["funding_ru"]))
    _add_heading(document, "Список литературы")
    _add_numbered_references(document, _references_ru())

    RU_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(RU_OUTPUT)


def _build_english_manuscript(
    *,
    profile: dict[str, object],
    regression_rows: list[dict[str, str]],
    classification_rows: list[dict[str, str]],
) -> None:
    document = Document()
    _set_default_font(document, "Times New Roman", 11)

    _add_centered_title(
        document,
        str(profile["title_en"]),
    )
    _add_centered_text(
        document,
        [
            format_author_names(profile, language="en"),
            *format_affiliations(profile, language="en"),
            format_emails(profile),
        ],
    )

    _add_heading(document, "Abstract")
    _add_paragraph(
        document,
        "This paper presents RuFLEX, a platform for hybrid deep fuzzy learning oriented toward geoanalytics, "
        "spatially informed classification, forecasting, and intelligent decision support. The platform combines "
        "interpretable rule-based mechanisms, differentiable fuzzy layers, and deep fuzzy feature-learning "
        "architectures within a unified workflow implemented in Python and PyTorch. Experimental validation was "
        "carried out on California Housing Regression and California Value Binary Geo."
    )
    _add_paragraph(
        document,
        f"Keywords: {format_keywords(profile['keywords_en'])}",
        italic_prefix="Keywords:",
    )

    _add_heading(document, "I. Introduction")
    _add_paragraph(
        document,
        "RuFLEX is designed as a platform rather than a single fixed architecture. It extends a vendored deep fuzzy "
        "backend with its own SDK, project model, explainability tooling, visual workbench, and reproducible "
        "experiment layer. This separation enables both backend continuity and independent platform evolution."
    )
    _add_paragraph(
        document,
        "The main objective is to support multiple hybrid deep fuzzy modeling modes while preserving interpretable "
        "reasoning at the level of membership functions, rules, hidden concepts, and factor contributions."
    )

    _add_heading(document, "II. Platform Architecture and Modeling Modes")
    _add_paragraph(
        document,
        "The current architecture includes a core layer with variables, terms, membership functions, rules, and "
        "model specifications; a model layer with flat neuro-fuzzy baselines and deep fuzzy feature learning; a "
        "training layer with bootstrap initialization, stage-wise pretraining, refinement, and fine-tuning; and an "
        "explainability layer with rule records, concept flow, and dashboard payloads."
    )
    _add_bullets(
        document,
        [
            "Python SDK and toolbox-style API",
            "visual workbench",
            "differentiable fuzzy training contour",
            "Explainability Dashboard",
            "reproducible article benchmark pipeline",
        ],
    )

    _add_heading(document, "III. Experimental Setup")
    _add_paragraph(
        document,
        "Two geo-oriented tabular tasks were used in validation: California Housing Regression and California Value "
        "Binary Geo. The first task preserves the spatial predictors Latitude and Longitude, while the second task "
        "uses the same source data to define a balanced binary classification problem based on housing value."
    )
    _add_paragraph(
        document,
        "The compared configurations were Flat Baseline, Flat Interpretable, Deep Article Demo, and Deep Research."
    )

    _add_heading(document, "IV. Results and Discussion")
    _add_paragraph(
        document,
        "On California Housing Regression, the best configuration was Deep Research with `test_rmse = 0.795757` "
        "and `test_r2 = 0.507159`. On California Value Binary Geo, the best configuration was Flat Interpretable "
        "with `test_accuracy = 0.746341` and `test_f1 = 0.739348`. This supports the platform-oriented design: "
        "different fuzzy regimes may become preferable depending on task structure."
    )

    _add_metric_table(
        document,
        title="Table I. Regression benchmark results",
        headers=("Configuration", "Test RMSE", "Test MAE", "Test R2"),
        rows=[
            (
                row["variant_label"],
                _fmt(row["test_rmse"]),
                _fmt(row["test_mae"]),
                _fmt(row["test_r2"]),
            )
            for row in regression_rows
        ],
    )
    _add_metric_table(
        document,
        title="Table II. Classification benchmark results",
        headers=("Configuration", "Test Accuracy", "Test Precision", "Test F1"),
        rows=[
            (
                row["variant_label"],
                _fmt(row["test_accuracy"]),
                _fmt(row["test_precision"]),
                _fmt(row["test_f1"]),
            )
            for row in classification_rows
        ],
    )

    _add_heading(document, "V. Explainability Figures")
    _add_figure(
        document,
        FIGURES_DIR / "figure1_regression_benchmark.png",
        "Fig. 1. Comparison of flat and deep fuzzy configurations on California Housing Regression.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure2_regression_board.png",
        "Fig. 2. Explainability board for the best regression configuration.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure3_classification_benchmark.png",
        "Fig. 3. Comparison of RuFLEX configurations on California Value Binary Geo.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure4_classification_board.png",
        "Fig. 4. Explainability board for the best classification configuration.",
        width=6.0,
    )
    _add_figure(
        document,
        FIGURES_DIR / "figure5_membership_example.png",
        "Fig. 5. Example membership functions used for interpretable fuzzification.",
        width=5.7,
    )

    _add_heading(document, "VI. Conclusion")
    _add_paragraph(
        document,
        "RuFLEX combines interpretable rule-based mechanisms, differentiable fuzzy layers, deep fuzzy feature "
        "learning, a Python SDK, a visual workbench, and explainability tooling within one reproducible workflow. "
        "The current implementation is suitable for both regression and classification scenarios and provides a "
        "foundation for future work on spatial fuzzy modules, stronger logical traceability, and counterfactual "
        "analysis."
    )
    _add_heading(document, "Acknowledgment")
    _add_paragraph(document, str(profile["acknowledgment_en"]))
    _add_heading(document, "Funding")
    _add_paragraph(document, str(profile["funding_en"]))
    _add_heading(document, "References")
    _add_numbered_references(document, _references_en())

    EN_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(EN_OUTPUT)


def _metric_title_paragraph(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True


def _add_metric_table(document: Document, *, title: str, headers: tuple[str, ...], rows: Iterable[tuple[str, ...]]) -> None:
    _metric_title_paragraph(document, title)
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_paragraph()


def _add_figure(document: Document, image_path: Path, caption: str, *, width: float) -> None:
    document.add_picture(str(image_path), width=Inches(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.add_run(caption)


def _add_numbered_references(document: Document, references: Iterable[str]) -> None:
    for item in references:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def _add_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True


def _add_centered_title(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(14)


def _add_centered_text(document: Document, lines: Iterable[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(line)


def _add_paragraph(document: Document, text: str, *, italic_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if italic_prefix is not None:
        prefix_run = paragraph.add_run(italic_prefix + " ")
        prefix_run.italic = True
        text = text.removeprefix(italic_prefix).strip()
    paragraph.add_run(text)


def _add_bullets(document: Document, items: Iterable[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _set_default_font(document: Document, font_name: str, size_pt: int) -> None:
    styles = document.styles
    for style_name in ("Normal",):
        style = styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(size_pt)


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _fmt(value: str) -> str:
    try:
        return f"{float(value):.6f}"
    except ValueError:
        return value


def _references_ru() -> list[str]:
    return [
        "Jang J.-S. R. ANFIS: Adaptive-Network-Based Fuzzy Inference System // IEEE Transactions on Systems, Man, and Cybernetics. 1993. Vol. 23, no. 3. P. 665-685.",
        "Paszke A., Gross S., Massa F. et al. PyTorch: An Imperative Style, High-Performance Deep Learning Library // Advances in Neural Information Processing Systems. 2019. Vol. 32.",
        "Pedregosa F., Varoquaux G., Gramfort A. et al. Scikit-learn: Machine Learning in Python // Journal of Machine Learning Research. 2011. Vol. 12. P. 2825-2830.",
        "Pace R. K., Barry R. Sparse Spatial Autoregressions // Statistics & Probability Letters. 1997. Vol. 33, no. 3. P. 291-297.",
        "Rudin C. Stop Explaining Black Box Machine Learning Models for High Stakes Decisions and Use Interpretable Models Instead // Nature Machine Intelligence. 2019. Vol. 1. P. 206-215.",
        "Ma X., Chen L., Deng Z. et al. Deep Image Feature Learning With Fuzzy Rules // IEEE Transactions on Emerging Topics in Computational Intelligence. 2024. Vol. 8. P. 724-737.",
        "Lebedeffson. deep-neuro-fuzzy [Электронный ресурс]. URL: https://github.com/lebedeffson/deep-neuro-fuzzy (дата обращения: 13.04.2026).",
    ]


def _references_en() -> list[str]:
    return [
        "J.-S. R. Jang, “ANFIS: Adaptive-Network-Based Fuzzy Inference System,” IEEE Transactions on Systems, Man, and Cybernetics, vol. 23, no. 3, pp. 665-685, 1993.",
        "A. Paszke, S. Gross, F. Massa et al., “PyTorch: An Imperative Style, High-Performance Deep Learning Library,” in Advances in Neural Information Processing Systems, vol. 32, 2019.",
        "F. Pedregosa, G. Varoquaux, A. Gramfort et al., “Scikit-learn: Machine Learning in Python,” Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011.",
        "R. K. Pace and R. Barry, “Sparse Spatial Autoregressions,” Statistics & Probability Letters, vol. 33, no. 3, pp. 291-297, 1997.",
        "C. Rudin, “Stop Explaining Black Box Machine Learning Models for High Stakes Decisions and Use Interpretable Models Instead,” Nature Machine Intelligence, vol. 1, pp. 206-215, 2019.",
        "X. Ma, L. Chen, Z. Deng et al., “Deep Image Feature Learning With Fuzzy Rules,” IEEE Transactions on Emerging Topics in Computational Intelligence, vol. 8, pp. 724-737, 2024.",
        "Lebedeffson, “deep-neuro-fuzzy,” GitHub repository. Available: https://github.com/lebedeffson/deep-neuro-fuzzy. Accessed: Apr. 13, 2026.",
    ]


if __name__ == "__main__":
    main()
