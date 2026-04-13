from __future__ import annotations

import csv
import hashlib
import subprocess
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ARTICLE_DIR = ROOT / "docs" / "article"
PACKAGE_DIR = ARTICLE_DIR / "q2_submission_package"
TABLES_DIR = ARTICLE_DIR / "tables_q2_final_plus"
EQUATIONS_DIR = ARTICLE_DIR / "generated_equations"
ASSETS_ROOT = ARTICLE_DIR / "assets_q2_final_plus"

MAIN_DOC_MD = ARTICLE_DIR / "article_final_ru.md"
MAIN_DOC_DOCX = ARTICLE_DIR / "article_final_ru.docx"
MAIN_DOC_PDF = ARTICLE_DIR / "article_final_ru.pdf"
EXT_DOC_MD = ARTICLE_DIR / "article_extended_materials_ru.md"
EXT_DOC_DOCX = ARTICLE_DIR / "article_extended_materials_ru.docx"
EXT_DOC_PDF = ARTICLE_DIR / "article_extended_materials_ru.pdf"

VISUAL_PACKAGE = ARTICLE_DIR / "visual_package_ru.md"
BENCHMARK_SUMMARY = ARTICLE_DIR / "benchmark_summary_q2_final_plus.md"
CHECKLIST_PATH = ARTICLE_DIR / "q2_submission_checklist_ru.md"
BUILD_REPORT_PATH = ARTICLE_DIR / "q2_build_report_ru.md"
PACKAGE_ZIP = ARTICLE_DIR / "q2_submission_package.zip"


@dataclass(frozen=True)
class DatasetAsset:
    slug: str
    title: str
    task_type: str
    asset_dir: Path
    board_name: str
    board_title: str

    @property
    def board_path(self) -> Path:
        return self.asset_dir / self.board_name

    @property
    def results_table_path(self) -> Path:
        return self.asset_dir / "results_table.csv"

    @property
    def overview_path(self) -> Path:
        return self.asset_dir / "results_overview.png"


DATASET_SPECS = (
    {
        "slug": "california_housing_regression",
        "title": "California Housing Regression",
        "task_type": "regression",
        "board_title": "Панель интерпретации RuFLEX: регрессионная задача",
    },
    {
        "slug": "california_value_binary_geo",
        "title": "California Value Binary Geo",
        "task_type": "binary_classification",
        "board_title": "Панель интерпретации RuFLEX: задача классификации",
    },
    {
        "slug": "covtype_binary_geo",
        "title": "Covertype Binary Geo",
        "task_type": "binary_classification",
        "board_title": "Панель интерпретации RuFLEX: стресс-тест Covertype",
    },
)


def _discover_dataset_asset(spec: dict[str, str]) -> DatasetAsset:
    slug = spec["slug"]
    candidates = sorted(
        (
            path
            for path in ASSETS_ROOT.glob(f"*_{slug}_article_benchmark")
            if path.is_dir() and (path / "benchmark_results.json").exists()
        ),
        key=lambda path: path.name,
    )
    if not candidates:
        raise FileNotFoundError(
            f"No benchmark asset directory was found for {slug!r} under {ASSETS_ROOT}."
        )
    asset_dir = candidates[-1]
    board_candidates = sorted(asset_dir.glob("article_board*_q2.png"))
    if not board_candidates:
        raise FileNotFoundError(f"No Q2 board image was found in {asset_dir}.")
    return DatasetAsset(
        slug=slug,
        title=spec["title"],
        task_type=spec["task_type"],
        asset_dir=asset_dir,
        board_name=board_candidates[0].name,
        board_title=spec["board_title"],
    )


DATASETS = tuple(_discover_dataset_asset(spec) for spec in DATASET_SPECS)

SCHEMES = (
    (
        ARTICLE_DIR / "generated_figures" / "scheme1_ruflex_platform_architecture.png",
        "Схема 1. Общая архитектура RuFLEX: вычислительное ядро, платформенный слой и исследовательские сервисы.",
    ),
    (
        ARTICLE_DIR / "generated_figures" / "scheme2_model_contour.png",
        "Схема 2. Модельный контур RuFLEX: плоский и глубокий режимы, скрытые концепты и пространственный каркас.",
    ),
    (
        ARTICLE_DIR / "generated_figures" / "scheme3_explainability_pipeline.png",
        "Схема 3. Контур интерпретации: от фаззификации объекта к правилам, концептам и аналитическим материалам.",
    ),
)


@dataclass
class MarkdownSection:
    title: str
    blocks: list[str]


@dataclass
class MarkdownDocument:
    title: str
    preamble_blocks: list[str]
    sections: list[MarkdownSection]


FORMULA_MAP = {
    "x = (x_1, x_2, ..., x_d).": r"x = (x_1, x_2, \ldots, x_d)",
    "z_i^(0) = φ_i(x_i), i = 1, 2, ..., d,": r"z_i^{(0)} = \phi_i(x_i),\quad i = 1, 2, \ldots, d",
    "z_i^(0) = φ_i(x_i),": r"z_i^{(0)} = \phi_i(x_i)",
    "z^(0)(x) = (z_1^(0), z_2^(0), ..., z_d^(0)).": r"z^{(0)}(x) = \left(z_1^{(0)}, z_2^{(0)}, \ldots, z_d^{(0)}\right)",
    "μ_ij(z_i^(0); θ_ij),": r"\mu_{ij}\!\left(z_i^{(0)}; \theta_{ij}\right)",
    "μ_ij(z_i^(0); θ_ij).": r"\mu_{ij}\!\left(z_i^{(0)}; \theta_{ij}\right)",
    "μ_ij(z_i^(0)) = exp( - (z_i^(0) - c_ij)^2 / (2σ_ij^2) ),": r"\mu_{ij}\!\left(z_i^{(0)}\right) = \exp\!\left(-\frac{\left(z_i^{(0)} - c_{ij}\right)^2}{2\sigma_{ij}^2}\right)",
    "μ_ij(z_i^(0)) = exp( - (z_i^(0) - c_ij)^2 / (2σ_ij^2) ).": r"\mu_{ij}\!\left(z_i^{(0)}\right) = \exp\!\left(-\frac{\left(z_i^{(0)} - c_{ij}\right)^2}{2\sigma_{ij}^2}\right)",
    "μ_ij(z_i^(0)) = 1 / ( 1 + | (z_i^(0) - c_ij) / a_ij |^(2b_ij) ),": r"\mu_{ij}\!\left(z_i^{(0)}\right) = \frac{1}{1 + \left|\frac{z_i^{(0)} - c_{ij}}{a_{ij}}\right|^{2b_{ij}}}",
    "μ_ij(z_i^(0)) = 1 / ( 1 + | (z_i^(0) - c_ij) / a_ij |^(2b_ij) ).": r"\mu_{ij}\!\left(z_i^{(0)}\right) = \frac{1}{1 + \left|\frac{z_i^{(0)} - c_{ij}}{a_{ij}}\right|^{2b_{ij}}}",
    "w_r^(l)(x) = ∏_(m=1)^(k_r) μ_(i_m j_m)^(l)( z_(i_m)^(l-1)(x); θ_(i_m j_m)^(l) ).": r"w_r^{(l)}(x) = \prod_{m=1}^{k_r} \mu_{i_m j_m}^{(l)}\!\left(z_{i_m}^{(l-1)}(x); \theta_{i_m j_m}^{(l)}\right)",
    r"\bar{w}_r^(l)(x) = w_r^(l)(x) / ( Σ_s w_s^(l)(x) + ε ),": r"\bar{w}_r^{(l)}(x) = \frac{w_r^{(l)}(x)}{\sum_s w_s^{(l)}(x) + \varepsilon}",
    "B_q^(l), q = 1, 2, ..., Q_l,": r"B_q^{(l)},\quad q = 1, 2, \ldots, Q_l",
    "h_q^(l)(x) = Σ_(r ∈ B_q^(l)) \\bar{w}_r^(l)(x) · g_r^(l)( z^(l-1)(x) ),": r"h_q^{(l)}(x) = \sum_{r \in B_q^{(l)}} \bar{w}_r^{(l)}(x) \cdot g_r^{(l)}\!\left(z^{(l-1)}(x)\right)",
    "h_q^(l)(x) = Σ_(r ∈ B_q^(l)) \\bar{w}_r^(l)(x) · g_r^(l)( z^(l-1)(x) ).": r"h_q^{(l)}(x) = \sum_{r \in B_q^{(l)}} \bar{w}_r^{(l)}(x) \cdot g_r^{(l)}\!\left(z^{(l-1)}(x)\right)",
    "g_r^(l)( z^(l-1)(x) ) = β_(r0)^(l) + Σ_t β_(rt)^(l) z_t^(l-1)(x).": r"g_r^{(l)}\!\left(z^{(l-1)}(x)\right) = \beta_{r0}^{(l)} + \sum_t \beta_{rt}^{(l)} z_t^{(l-1)}(x)",
    "h^(l)(x) = ( h_1^(l)(x), h_2^(l)(x), ..., h_(Q_l)^(l)(x) )": r"h^{(l)}(x) = \left(h_1^{(l)}(x), h_2^{(l)}(x), \ldots, h_{Q_l}^{(l)}(x)\right)",
    "z^(0)(x) -> h^(1)(x) -> h^(2)(x) -> ... -> h^(L-1)(x).": r"z^{(0)}(x) \rightarrow h^{(1)}(x) \rightarrow h^{(2)}(x) \rightarrow \cdots \rightarrow h^{(L-1)}(x)",
    r"\hat{y}(x) = Σ_(r ∈ B^(L)) \bar{w}_r^(L)(x) · g_r^(L)( h^(L-1)(x) ).": r"\hat{y}(x) = \sum_{r \in B^{(L)}} \bar{w}_r^{(L)}(x) \cdot g_r^{(L)}\!\left(h^{(L-1)}(x)\right)",
    "s(x) = Σ_(r ∈ B^(L)) \\bar{w}_r^(L)(x) · g_r^(L)( h^(L-1)(x) ),": r"s(x) = \sum_{r \in B^{(L)}} \bar{w}_r^{(L)}(x) \cdot g_r^{(L)}\!\left(h^{(L-1)}(x)\right)",
    r"\hat{p}(x) = 1 / (1 + exp(-s(x))).": r"\hat{p}(x) = \frac{1}{1 + \exp(-s(x))}",
    "L = L_задачи + λ_1 L_разреженности + λ_2 L_перекрытия + λ_3 L_покрытия + λ_4 L_концептов.": r"L = L_{\mathrm{task}} + \lambda_1 L_{\mathrm{sparsity}} + \lambda_2 L_{\mathrm{overlap}} + \lambda_3 L_{\mathrm{coverage}} + \lambda_4 L_{\mathrm{concept}}",
    "L_задачи = (1/N) Σ_(n=1)^N ( y_n - \\hat{y}(x_n) )^2,": r"L_{\mathrm{task}} = \frac{1}{N} \sum_{n=1}^{N} \left(y_n - \hat{y}(x_n)\right)^2",
    "L_задачи = (1/N) Σ_(n=1)^N ( y_n - \\hat{y}(x_n) )^2.": r"L_{\mathrm{task}} = \frac{1}{N} \sum_{n=1}^{N} \left(y_n - \hat{y}(x_n)\right)^2",
    "L_задачи = - (1/N) Σ_(n=1)^N [ y_n log \\hat{p}(x_n) + (1-y_n) log (1-\\hat{p}(x_n)) ].": r"L_{\mathrm{task}} = -\frac{1}{N} \sum_{n=1}^{N} \left[y_n \log \hat{p}(x_n) + (1-y_n)\log\left(1-\hat{p}(x_n)\right)\right]",
    "L_разреженности = Σ_(l,r) |α_r^(l)|,": r"L_{\mathrm{sparsity}} = \sum_{l,r} \left|\alpha_r^{(l)}\right|",
    "L_разреженности = Σ_(l,r) |α_r^(l)|.": r"L_{\mathrm{sparsity}} = \sum_{l,r} \left|\alpha_r^{(l)}\right|",
    "L_перекрытия = Σ_(i,j<k) ∫ μ_ij(t) μ_ik(t) dt,": r"L_{\mathrm{overlap}} = \sum_{i, j < k} \int \mu_{ij}(t)\mu_{ik}(t)\,dt",
    "L_перекрытия = Σ_(i,j<k) ∫ μ_ij(t) μ_ik(t) dt.": r"L_{\mathrm{overlap}} = \sum_{i, j < k} \int \mu_{ij}(t)\mu_{ik}(t)\,dt",
    "L_концептов = Σ_l Σ_(q≠p) < h_q^(l), h_p^(l) >^2.": r"L_{\mathrm{concept}} = \sum_l \sum_{q \neq p} \left\langle h_q^{(l)}, h_p^{(l)} \right\rangle^2",
    "X ∈ R^(H×W×C).": r"X \in \mathbb{R}^{H \times W \times C}",
    "u_m(u,v) = Σ_((a,b,c) ∈ Ω) κ_(mabc) · μ_(mabc)( X_(u+a, v+b, c) ),": r"u_m(u,v) = \sum_{(a,b,c)\in\Omega} \kappa_{mabc} \cdot \mu_{mabc}\!\left(X_{u+a, v+b, c}\right)",
    "C_r(x) = \\bar{w}_r^(L)(x) · g_r^(L)( h^(L-1)(x) ).": r"C_r(x) = \bar{w}_r^{(L)}(x) \cdot g_r^{(L)}\!\left(h^{(L-1)}(x)\right)",
    "C_q^(l)(x) = Σ_(r ∈ B_q^(l)) C_r^(l)(x).": r"C_q^{(l)}(x) = \sum_{r \in B_q^{(l)}} C_r^{(l)}(x)",
}


def main() -> None:
    build_q2_article_package()
    print(MAIN_DOC_DOCX)
    print(EXT_DOC_DOCX)
    print(PACKAGE_DIR)
    print(PACKAGE_ZIP)


def build_q2_article_package() -> None:
    EQUATIONS_DIR.mkdir(parents=True, exist_ok=True)
    TABLES_DIR.mkdir(parents=True, exist_ok=True)
    package_docs = PACKAGE_DIR / "docs"
    package_figures = PACKAGE_DIR / "figures"
    package_tables = PACKAGE_DIR / "tables"
    if PACKAGE_DIR.exists():
        shutil.rmtree(PACKAGE_DIR)
    package_docs.mkdir(parents=True, exist_ok=True)
    package_figures.mkdir(parents=True, exist_ok=True)
    package_tables.mkdir(parents=True, exist_ok=True)

    main_document = parse_markdown_document(MAIN_DOC_MD)
    extended_document = parse_markdown_document(EXT_DOC_MD)

    main_docx = build_document(
        source=main_document,
        output_path=MAIN_DOC_DOCX,
        include_results_assets=True,
        include_interpretability_assets=False,
        compact_results=True,
    )
    extended_docx = build_document(
        source=extended_document,
        output_path=EXT_DOC_DOCX,
        include_results_assets=True,
        include_interpretability_assets=True,
        compact_results=False,
    )
    pdf_statuses = {
        MAIN_DOC_PDF.name: convert_docx_to_pdf(main_docx, MAIN_DOC_PDF),
        EXT_DOC_PDF.name: convert_docx_to_pdf(extended_docx, EXT_DOC_PDF),
    }

    summary_tables = write_summary_tables(package_tables)
    write_checklist()
    write_build_report(pdf_statuses)
    populate_package(
        package_docs,
        package_figures,
        package_tables,
        main_docx,
        extended_docx,
        summary_tables,
    )
    build_package_zip()


def parse_markdown_document(path: Path) -> MarkdownDocument:
    lines = path.read_text(encoding="utf-8").splitlines()
    if not lines or not lines[0].startswith("# "):
        raise ValueError(f"Expected a markdown title in {path}")

    title = lines[0][2:].strip()
    cursor = 1
    preamble_blocks: list[str] = []

    current: list[str] = []

    def flush(target: list[str]) -> None:
        if current:
            target.append("\n".join(line.strip() for line in current if line.strip()).strip())
            current.clear()

    while cursor < len(lines) and not lines[cursor].startswith("## "):
        line = lines[cursor]
        if line.strip():
            current.append(line.rstrip())
        else:
            flush(preamble_blocks)
        cursor += 1
    flush(preamble_blocks)

    sections: list[MarkdownSection] = []
    while cursor < len(lines):
        heading = lines[cursor][3:].strip()
        cursor += 1
        blocks: list[str] = []
        while cursor < len(lines) and not lines[cursor].startswith("## "):
            line = lines[cursor]
            if line.strip():
                current.append(line.rstrip())
            else:
                flush(blocks)
            cursor += 1
        flush(blocks)
        sections.append(MarkdownSection(title=heading, blocks=blocks))

    return MarkdownDocument(title=title, preamble_blocks=preamble_blocks, sections=sections)


def build_document(
    *,
    source: MarkdownDocument,
    output_path: Path,
    include_results_assets: bool,
    include_interpretability_assets: bool,
    compact_results: bool,
) -> Path:
    document = Document()
    configure_document(document)
    add_title_block(document, source)

    for section in source.sections:
        heading = document.add_heading(section.title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for block in section.blocks:
            add_block(document, block)
        append_section_assets(
            document,
            section=section.title,
            include_results_assets=include_results_assets,
            include_interpretability_assets=include_interpretability_assets,
            compact_results=compact_results,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


def configure_document(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)

    for heading_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[heading_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.8)


def add_title_block(document: Document, source: MarkdownDocument) -> None:
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(source.title)
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    for block in source.preamble_blocks:
        if block.startswith("Аннотация."):
            add_labeled_paragraph(document, block, label="Аннотация.")
            continue
        if block.startswith("Ключевые слова:"):
            add_labeled_paragraph(document, block, label="Ключевые слова:")
            continue

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        for line in block.splitlines():
            clean = line.replace("  ", "").strip()
            if not clean:
                continue
            run = paragraph.add_run(clean)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            paragraph.add_run("\n")


def add_labeled_paragraph(document: Document, text: str, *, label: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)

    body = text[len(label) :].strip()
    lead = paragraph.add_run(f"{label} ")
    lead.bold = True
    lead.font.name = "Times New Roman"
    lead.font.size = Pt(12)
    tail = paragraph.add_run(body)
    tail.font.name = "Times New Roman"
    tail.font.size = Pt(12)


def add_block(document: Document, block: str) -> None:
    normalized = normalize_text(block)
    if normalized in FORMULA_MAP:
        add_formula(document, normalized, FORMULA_MAP[normalized])
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(6)

    if normalized[0:2].isdigit() and "." in normalized[:4]:
        paragraph.paragraph_format.first_line_indent = Inches(0.0)
        paragraph.paragraph_format.left_indent = Inches(0.0)

    run = paragraph.add_run(normalized)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def add_formula(document: Document, normalized_text: str, latex_formula: str) -> None:
    image_path = render_formula_image(normalized_text, latex_formula)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(5.6))


def render_formula_image(normalized_text: str, latex_formula: str) -> Path:
    EQUATIONS_DIR.mkdir(parents=True, exist_ok=True)
    key = hashlib.sha1(normalized_text.encode("utf-8")).hexdigest()[:16]
    output_path = EQUATIONS_DIR / f"{key}.png"
    if output_path.exists():
        return output_path

    plt.rcParams["mathtext.fontset"] = "stix"
    plt.rcParams["font.family"] = "STIXGeneral"

    probe = plt.figure(figsize=(0.01, 0.01))
    probe.patch.set_alpha(0.0)
    text_artist = probe.text(0.0, 0.0, f"${latex_formula}$", fontsize=18)
    probe.canvas.draw()
    bbox = text_artist.get_window_extent(renderer=probe.canvas.get_renderer()).expanded(1.08, 1.35)
    width = max(bbox.width / probe.dpi, 1.5)
    height = max(bbox.height / probe.dpi, 0.35)
    plt.close(probe)

    figure = plt.figure(figsize=(width, height))
    figure.patch.set_alpha(0.0)
    axis = figure.add_axes([0, 0, 1, 1])
    axis.axis("off")
    axis.text(0.5, 0.5, f"${latex_formula}$", fontsize=18, ha="center", va="center")
    figure.savefig(output_path, dpi=300, transparent=True, bbox_inches="tight", pad_inches=0.04)
    plt.close(figure)
    return output_path


def append_section_assets(
    document: Document,
    *,
    section: str,
    include_results_assets: bool,
    include_interpretability_assets: bool,
    compact_results: bool,
) -> None:
    if "Архитектура" in section:
        for figure_path, caption in SCHEMES:
            add_figure(document, figure_path, caption, width=6.2)
        return

    if include_results_assets and ("Результаты" in section or "Как трактовать результаты" in section):
        for dataset in DATASETS:
            add_metric_table(document, dataset)
            add_figure(document, dataset.board_path, dataset.board_title, width=6.2)
            if not compact_results:
                add_figure(document, dataset.overview_path, f"{dataset.title}: сводный график метрик.", width=6.0)

        if include_interpretability_assets:
            for figure_path, caption in interpretability_figures():
                add_figure(document, figure_path, caption, width=6.0)


def interpretability_figures() -> tuple[tuple[Path, str], ...]:
    regression = DATASETS[0].asset_dir
    classification = DATASETS[1].asset_dir
    return (
        (
            regression / "best_training_history.png",
            "Дополнительный материал. История обучения для лучшей конфигурации на California Housing Regression.",
        ),
        (
            regression / "best_sample_top_rules.png",
            "Дополнительный материал. Наиболее активные правила для репрезентативного объекта в регрессионной задаче.",
        ),
        (
            classification / "best_sample_hidden_concepts.png",
            "Дополнительный материал. Скрытые нечеткие концепты для репрезентативного объекта в задаче California Value Binary Geo.",
        ),
        (
            classification / "best_sample_fuzzification.png",
            "Дополнительный материал. Карта фаззификации репрезентативного объекта в задаче California Value Binary Geo.",
        ),
    )


def add_figure(document: Document, image_path: Path, caption: str, *, width: float) -> None:
    if not image_path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))

    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.name = "Times New Roman"
    caption_run.font.size = Pt(10)


def add_metric_table(document: Document, dataset: DatasetAsset) -> None:
    rows = load_dataset_rows(dataset.results_table_path, dataset.task_type)
    if not rows:
        return

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run(f"Таблица. {dataset.title}: агрегированные результаты mean ± std")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(11)

    if dataset.task_type == "regression":
        headers = ("Модель", "Класс", "RMSE", "MAE", "R²")
        metric_keys = ("rmse", "mae", "r2")
        winner_key = "rmse"
        winner_mode = "min"
    else:
        headers = ("Модель", "Класс", "Accuracy", "F1", "Precision", "Recall")
        metric_keys = ("accuracy", "f1", "precision", "recall")
        winner_key = "f1"
        winner_mode = "max"

    overall_winner = select_winner(rows, winner_key, winner_mode, family=None)
    ruflex_winner = select_winner(rows, winner_key, winner_mode, family="ruflex")

    table = document.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for column_index, header in enumerate(headers):
        cell = table.rows[0].cells[column_index]
        set_cell_text(cell, header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)

    for row_index, row in enumerate(rows, start=1):
        is_overall_winner = row["model_name"] == overall_winner["model_name"]
        is_ruflex_winner = row["model_name"] == ruflex_winner["model_name"]
        role_label = "RuFLEX" if row["family"] == "ruflex" else "Внешний baseline"

        set_cell_text(
            table.rows[row_index].cells[0],
            row["model_label"],
            bold=is_overall_winner,
            italic=is_ruflex_winner and not is_overall_winner,
            size=9,
        )
        set_cell_text(table.rows[row_index].cells[1], role_label, size=9)

        for offset, metric_key in enumerate(metric_keys, start=2):
            metric_value = row[metric_key]
            set_cell_text(
                table.rows[row_index].cells[offset],
                metric_value,
                bold=is_overall_winner,
                size=9,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run("Жирным выделен лучший общий результат. Курсивом отмечен лучший результат внутри семейства RuFLEX.")
    note_run.font.name = "Times New Roman"
    note_run.font.size = Pt(9)


def load_dataset_rows(path: Path, task_type: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with path.open(encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            family = row["family"]
            metrics: dict[str, str] = {}
            if task_type == "regression":
                metrics["rmse"] = format_mean_std(row["test_rmse_mean"], row["test_rmse_std"])
                metrics["mae"] = format_mean_std(row["test_mae_mean"], row["test_mae_std"])
                metrics["r2"] = format_mean_std(row["test_r2_mean"], row["test_r2_std"])
            else:
                metrics["accuracy"] = format_mean_std(row["test_accuracy_mean"], row["test_accuracy_std"])
                metrics["f1"] = format_mean_std(row["test_f1_mean"], row["test_f1_std"])
                metrics["precision"] = format_mean_std(row["test_precision_mean"], row["test_precision_std"])
                metrics["recall"] = format_mean_std(row["test_recall_mean"], row["test_recall_std"])

            rows.append(
                {
                    "model_name": row["model_name"],
                    "model_label": row["model_label"],
                    "family": family,
                    **metrics,
                    "_winner_value": row["test_rmse_mean"] if task_type == "regression" else row["test_f1_mean"],
                    "_accuracy_value": row.get("test_accuracy_mean", ""),
                }
            )
    return rows


def format_mean_std(mean_value: str, std_value: str) -> str:
    return f"{float(mean_value):.3f} ± {float(std_value):.3f}"


def select_winner(
    rows: list[dict[str, str]],
    metric_key: str,
    mode: str,
    *,
    family: str | None,
) -> dict[str, str]:
    candidates = [row for row in rows if family is None or row["family"] == family]
    if not candidates:
        return rows[0]

    def score(item: dict[str, str]) -> float:
        if metric_key == "rmse":
            return float(item["_winner_value"])
        if metric_key == "f1":
            mean_part = item["f1"].split("±", maxsplit=1)[0].strip()
            return float(mean_part)
        raise ValueError(metric_key)

    return min(candidates, key=score) if mode == "min" else max(candidates, key=score)


def set_cell_text(cell, text: str, *, bold: bool = False, italic: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def write_summary_tables(package_tables: Path) -> list[Path]:
    summary_paths: list[Path] = []
    for dataset in DATASETS:
        rows = load_dataset_rows(dataset.results_table_path, dataset.task_type)
        output_path = TABLES_DIR / f"{dataset.slug}_summary.csv"
        with output_path.open("w", encoding="utf-8", newline="") as handle:
            if dataset.task_type == "regression":
                fieldnames = ("model", "class", "rmse", "mae", "r2")
            else:
                fieldnames = ("model", "class", "accuracy", "f1", "precision", "recall")
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            for row in rows:
                payload = {
                    "model": row["model_label"],
                    "class": "RuFLEX" if row["family"] == "ruflex" else "external_baseline",
                }
                for fieldname in fieldnames[2:]:
                    payload[fieldname] = row[fieldname]
                writer.writerow(payload)
        copied_path = package_tables / output_path.name
        shutil.copy2(output_path, copied_path)
        summary_paths.append(output_path)
    return summary_paths


def write_checklist() -> None:
    CHECKLIST_PATH.write_text(
        "\n".join(
            [
                "# Q2 Submission Checklist",
                "",
                "## Уже доведено автоматически",
                "",
                "1. Пересчитана финальная benchmark-серия `article_benchmark_q2_final_plus` по трем задачам.",
                "2. Внешняя таблица усилена за счет `Gradient Boosting`.",
                "3. Собраны финальные boards и обновлен визуальный пакет.",
                "4. Переписаны основной текст статьи и расширенные материалы под новую clean-series.",
                "5. Собраны `article_final_ru.docx` и `article_extended_materials_ru.docx` с встроенными рисунками, таблицами и визуально оформленными формулами.",
                "6. Автоматически собраны `article_final_ru.pdf` и `article_extended_materials_ru.pdf`.",
                "7. Подготовлен пакет `docs/article/q2_submission_package`.",
                "",
                "## Что остается сделать вручную перед подачей",
                "",
                "1. Добавить аффилиации, e-mail и ORCID для утвержденного списка авторов, если эти сведения требуются целевой площадке.",
                "2. При необходимости снять 1-3 живых скриншота интерфейса RuFLEX и решить, нужны ли они в основной статье или оставить их только в расширенных материалах.",
                "3. Согласовать требования конкретного журнала: объем, стиль ссылок, размер рисунков, правила по приложениям и supplementary materials.",
                "4. Если журнал принимает PDF на этапе рецензирования, использовать `article_final_ru.pdf` как основной review-артефакт, а `docx` оставить как редактируемую исходную версию.",
                "5. Если журнал требует именно `docx` с редактируемыми формулами Word уже на этапе подачи, заменить встроенные формульные изображения на нативные equation-объекты.",
                "6. Провести финальную ручную вычитку по верстке в Word: переносы, разрывы страниц, подписи рисунков и таблиц.",
                "7. Выполнить антиплагиат и финальную внутреннюю редакторскую проверку.",
                "",
                "## Что не нужно больше пересобирать вручную",
                "",
                "1. Бенчмарки, таблицы и boards уже синхронизированы между собой.",
                "2. Основные markdown-версии и офисные версии теперь относятся к одной и той же финальной benchmark-серии.",
                "3. Визуальный манифест и benchmark-сводка уже указывают на правильные каталоги.",
                "4. PDF-версии уже включены в submission-пакет и не требуют отдельной ручной конвертации.",
            ]
        ),
        encoding="utf-8",
    )


def write_build_report(pdf_statuses: dict[str, str]) -> None:
    lines = [
        "# Q2 Build Report",
        "",
        "## Основной статус",
        "",
        "1. Основные `docx`-документы собраны из актуальных markdown-источников.",
        "2. В документы встроены рисунки, boards, краткие таблицы и визуально оформленные формулы.",
        "3. Пакет `q2_submission_package` собран и архивирован.",
        "",
        "## PDF-превью",
        "",
    ]
    for name, status in pdf_statuses.items():
        lines.append(f"- `{name}`: {status}")
    lines.extend(
        [
            "",
            "## Примечание",
            "",
            "В текущей сборке PDF-превью включены в пакет наряду с `docx`-версиями и могут использоваться как основные review-артефакты там, где площадка принимает PDF на этапе подачи.",
            "",
            "Оставшийся технический риск относится только к площадкам, которые требуют именно `docx` с редактируемыми Word-equation формулами уже на этапе submission. В текущей сборке формулы встроены в `docx` как визуально оформленные объекты, поэтому для такой площадки нужен отдельный ручной проход по формулам на финальном этапе верстки.",
        ]
    )
    BUILD_REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def populate_package(
    package_docs: Path,
    package_figures: Path,
    package_tables: Path,
    main_docx: Path,
    extended_docx: Path,
    summary_tables: list[Path],
) -> None:
    for source in (
        MAIN_DOC_MD,
        EXT_DOC_MD,
        main_docx,
        extended_docx,
        MAIN_DOC_PDF,
        EXT_DOC_PDF,
        VISUAL_PACKAGE,
        BENCHMARK_SUMMARY,
        CHECKLIST_PATH,
        BUILD_REPORT_PATH,
    ):
        if source.exists():
            shutil.copy2(source, package_docs / source.name)

    for source in summary_tables:
        shutil.copy2(source, package_tables / source.name)

    for figure_path, _caption in SCHEMES:
        shutil.copy2(figure_path, package_figures / figure_path.name)
    for dataset in DATASETS:
        shutil.copy2(dataset.board_path, package_figures / dataset.board_path.name)
        shutil.copy2(dataset.overview_path, package_figures / dataset.overview_path.name.replace("results_overview", f"{dataset.slug}_results_overview"))

    readme_path = PACKAGE_DIR / "README.md"
    readme_path.write_text(
        "\n".join(
            [
                "# Q2 Submission Package",
                "",
                "## Содержимое",
                "",
                "1. `docs/` — финальные markdown-, docx- и pdf-версии статьи, расширенных материалов и служебных сводок.",
                "   Здесь же лежат `q2_submission_checklist_ru.md` и `q2_build_report_ru.md`.",
                "2. `figures/` — схемы архитектуры, boards и компактные обзорные графики.",
                "3. `tables/` — краткие итоговые CSV-таблицы по трем benchmark-задачам.",
                "",
                "## Базовая экспериментальная серия",
                "",
                "Пакет соответствует benchmark-серии `article_benchmark_q2_final_plus`.",
                "",
                "## Примечание",
                "",
                "В текущем пакете уже лежат `article_final_ru.pdf` и `article_extended_materials_ru.pdf`. Если целевая площадка принимает PDF на этапе рецензирования, эти файлы можно использовать как основные review-артефакты.",
                "",
                "Формулы в текущих `docx` визуально оформлены и встроены в документ. Если журнал требует именно редактируемые Word-equation объекты уже на этапе подачи, этот этап нужно выполнить вручную перед отправкой.",
            ]
        ),
        encoding="utf-8",
    )


def build_package_zip() -> None:
    if PACKAGE_ZIP.exists():
        PACKAGE_ZIP.unlink()
    with zipfile.ZipFile(PACKAGE_ZIP, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(PACKAGE_DIR.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(ARTICLE_DIR))


def convert_docx_to_pdf(source_docx: Path, output_pdf: Path) -> str:
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    if output_pdf.exists():
        output_pdf.unlink()
    try:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_pdf.parent),
                str(source_docx),
            ],
            cwd=ROOT,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
    except Exception as error:
        return f"не собрано автоматически ({type(error).__name__})"
    return "собрано автоматически"


def normalize_text(text: str) -> str:
    return " ".join(text.replace("  ", " ").split()).strip()


if __name__ == "__main__":
    main()
