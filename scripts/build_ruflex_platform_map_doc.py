from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
ARTICLE_DIR = ROOT / "docs" / "article"
SOURCE_MD = ARTICLE_DIR / "ruflex_platform_map_ru.md"
OUTPUT_DOCX = ARTICLE_DIR / "ruflex_platform_map_ru.docx"

INLINE_TOKEN_RE = re.compile(r"(`[^`]+`|\[[^\]]+\]\([^)]+\))")
ORDERED_ITEM_RE = re.compile(r"^\d+\.\s+")


def main() -> None:
    output_path = build_platform_map_docx()
    print(output_path)


def build_platform_map_docx() -> Path:
    document = Document()
    _configure_document(document)
    _render_markdown(document, SOURCE_MD.read_text(encoding="utf-8").splitlines())
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.0)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    for style_name, size in (
        ("Title", 18),
        ("Heading 1", 15),
        ("Heading 2", 13),
        ("Heading 3", 12),
    ):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)


def _render_markdown(document: Document, lines: list[str]) -> None:
    paragraph_lines: list[str] = []
    math_lines: list[str] = []
    in_math_block = False

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines if line.strip())
        paragraph_lines.clear()
        _add_paragraph(document, text, style="Normal")

    def flush_math() -> None:
        if not math_lines:
            return
        text = " ".join(line.strip() for line in math_lines if line.strip())
        math_lines.clear()
        paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name = "Cambria Math"
        run.font.size = Pt(11)
        _set_run_preserve_spaces(run)

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if in_math_block:
            if stripped == r"\]":
                flush_math()
                in_math_block = False
            else:
                math_lines.append(stripped)
            continue

        if stripped == r"\[":
            flush_paragraph()
            in_math_block = True
            math_lines.clear()
            continue

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            _add_heading(document, stripped[2:].strip(), level=0)
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            _add_heading(document, stripped[3:].strip(), level=1)
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            _add_heading(document, stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("#### "):
            flush_paragraph()
            _add_heading(document, stripped[5:].strip(), level=3)
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            _add_paragraph(document, stripped[2:].strip(), style="List Bullet")
            continue

        if ORDERED_ITEM_RE.match(stripped):
            flush_paragraph()
            item_text = ORDERED_ITEM_RE.sub("", stripped, count=1)
            _add_paragraph(document, item_text, style="List Number")
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()
    if in_math_block:
        flush_math()


def _add_heading(document: Document, text: str, *, level: int) -> None:
    if level == 0:
        paragraph = document.add_paragraph(style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _render_inline(paragraph, text)
        return
    paragraph = document.add_paragraph(style=f"Heading {level}")
    _render_inline(paragraph, text)


def _add_paragraph(document: Document, text: str, *, style: str) -> None:
    paragraph = document.add_paragraph(style=style)
    _render_inline(paragraph, text)


def _render_inline(paragraph, text: str) -> None:
    position = 0
    for match in INLINE_TOKEN_RE.finditer(text):
        if match.start() > position:
            _add_text_run(paragraph, text[position : match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
            _set_run_preserve_spaces(run)
        else:
            label, target = _parse_link(token)
            _add_text_run(paragraph, f"{label} ({target})")
        position = match.end()
    if position < len(text):
        _add_text_run(paragraph, text[position:])


def _add_text_run(paragraph, text: str) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    _set_run_preserve_spaces(run)


def _set_run_preserve_spaces(run) -> None:
    text_nodes = run._r.xpath(".//w:t")
    for node in text_nodes:
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _parse_link(token: str) -> tuple[str, str]:
    match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
    if not match:
        return token, ""
    return match.group(1), match.group(2)


if __name__ == "__main__":
    main()
