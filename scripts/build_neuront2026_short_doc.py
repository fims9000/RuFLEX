from __future__ import annotations

import csv
import subprocess
import tempfile
import urllib.request
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
ARTICLE_DIR = ROOT / "docs" / "article"
SOURCE_MD = ARTICLE_DIR / "neuront2026_short_ru.md"
OUTPUT_DOCX = ARTICLE_DIR / "neuront2026_short_ru.docx"
OUTPUT_DOC = ARTICLE_DIR / "neuront2026_short_ru.doc"
OUTPUT_PDF = ARTICLE_DIR / "neuront2026_short_ru.pdf"
TEMPLATE_URL = "https://neuront.etu.ru/assets/files/shablon-dokladov.doc"
DOCUMENT_AUTHORS = (
    "Лебедев М.Д.; Трофимов Ю.В.; Лебедев А.Д.; "
    "Аверкин А.Н.; Ильин А.С.; Алексеев А.К."
)
DOCUMENT_KEYWORDS = (
    "RuFLEX; геоданные; глубокое нечеткое обучение; "
    "нейро-нечеткие модели; интерпретируемость"
)
FIGURES = {
    "platform_architecture": ARTICLE_DIR / "generated_figures" / "scheme1_ruflex_platform_architecture.png",
}

RESULT_TABLES = {
    "california_housing_regression": ARTICLE_DIR
    / "tables_q2_final_plus"
    / "california_housing_regression_summary.csv",
    "california_value_binary_geo": ARTICLE_DIR
    / "tables_q2_final_plus"
    / "california_value_binary_geo_summary.csv",
    "covtype_binary_geo": ARTICLE_DIR
    / "tables_q2_final_plus"
    / "covtype_binary_geo_summary.csv",
}


def main() -> None:
    paths = build_neuront2026_short_package()
    for path in paths:
        print(path)


def build_neuront2026_short_package() -> tuple[Path, Path, Path]:
    with tempfile.TemporaryDirectory(prefix="neuront2026_short_") as temp_dir:
        temp_root = Path(temp_dir)
        template_doc = temp_root / "neuront_template.doc"
        template_docx = temp_root / "neuront_template.docx"
        _download_template(template_doc)
        _convert_to_docx(template_doc, temp_root)

        document = Document(template_docx)
        _clear_document_body(document)
        _populate_document(document)
        _populate_metadata(document)

        OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
        document.save(OUTPUT_DOCX)
        _convert_to_doc(OUTPUT_DOCX, OUTPUT_DOC.parent)
        _convert_to_pdf(OUTPUT_DOCX, OUTPUT_PDF.parent)
    return OUTPUT_DOCX, OUTPUT_DOC, OUTPUT_PDF


def _download_template(destination: Path) -> None:
    urllib.request.urlretrieve(TEMPLATE_URL, destination)


def _convert_to_docx(source_doc: Path, output_dir: Path) -> None:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "docx",
            "--outdir",
            str(output_dir),
            str(source_doc),
        ],
        check=True,
        cwd=ROOT,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def _convert_to_doc(source_docx: Path, output_dir: Path) -> None:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "doc:MS Word 97",
            "--outdir",
            str(output_dir),
            str(source_docx),
        ],
        check=True,
        cwd=ROOT,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def _convert_to_pdf(source_docx: Path, output_dir: Path) -> None:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(source_docx),
        ],
        check=True,
        cwd=ROOT,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def _clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def _populate_document(document: Document) -> None:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    title = lines[0][2:].strip()

    cursor = 1
    preamble_blocks: list[str] = []
    current: list[str] = []

    def flush(target: list[str]) -> None:
        if current:
            target.append("\n".join(line.strip() for line in current if line.strip()).strip())
            current.clear()

    while cursor < len(lines) and not lines[cursor].startswith("## "):
        line = lines[cursor]
        if line.strip():
            current.append(line.rstrip())
        else:
            flush(preamble_blocks)
        cursor += 1
    flush(preamble_blocks)

    sections: list[tuple[str, list[str]]] = []
    while cursor < len(lines):
        heading = lines[cursor][3:].strip()
        cursor += 1
        blocks: list[str] = []
        while cursor < len(lines) and not lines[cursor].startswith("## "):
            line = lines[cursor]
            if line.strip():
                current.append(line.rstrip())
            else:
                flush(blocks)
            cursor += 1
        flush(blocks)
        sections.append((heading, blocks))

    _add_paragraph(document, title, style="Заголовок статьи", align=WD_ALIGN_PARAGRAPH.CENTER)
    for index, block in enumerate(preamble_blocks):
        if index == 0:
            _add_paragraph(document, block, style="Автор", align=WD_ALIGN_PARAGRAPH.CENTER)
            continue
        if block.startswith("Аннотация."):
            _add_paragraph(document, block, style="Аннотация")
            continue
        if block.startswith("Ключевые слова:"):
            _add_paragraph(document, block, style="Ключевые слова")
            continue
        _add_paragraph(document, block, style="Организация", align=WD_ALIGN_PARAGRAPH.CENTER)

    for heading, blocks in sections:
        if heading == "Список литературы":
            _add_paragraph(document, heading, style="Heading 5")
            for block in blocks:
                _add_paragraph(document, block, style="список лит-ры")
            continue

        _add_paragraph(document, heading, style="Heading 1")
        for block in blocks:
            if block == "[[TABLE:main_results]]":
                _add_results_table(document)
            elif block == "[[FIGURE:platform_architecture]]":
                _add_figure(
                    document,
                    FIGURES["platform_architecture"],
                    "Рис. 1. Архитектура платформы RuFLEX.",
                )
            else:
                _add_paragraph(document, block, style="Body Text")


def _populate_metadata(document: Document) -> None:
    title = SOURCE_MD.read_text(encoding="utf-8").splitlines()[0][2:].strip()
    core = document.core_properties
    core.title = title
    core.author = DOCUMENT_AUTHORS
    core.subject = "Краткая версия статьи для NEURONT 2026"
    core.keywords = DOCUMENT_KEYWORDS
    core.comments = "Собрано автоматически по официальному шаблону NEURONT 2026."


def _add_paragraph(
    document: Document,
    text: str,
    *,
    style: str,
    align: WD_ALIGN_PARAGRAPH | None = None,
) -> None:
    paragraph = document.add_paragraph(style=style)
    if align is not None:
        paragraph.alignment = align
    paragraph.add_run(text)


def _add_figure(document: Document, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(7.3))

    caption_paragraph = document.add_paragraph(style="Body Text")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    caption_run.font.size = Pt(8)


def _read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _select_row(rows: list[dict[str, str]], model: str) -> dict[str, str]:
    for row in rows:
        if row["model"] == model:
            return row
    raise KeyError(model)


def _add_results_table(document: Document) -> None:
    reg_rows = _read_rows(RESULT_TABLES["california_housing_regression"])
    cls_rows = _read_rows(RESULT_TABLES["california_value_binary_geo"])
    cov_rows = _read_rows(RESULT_TABLES["covtype_binary_geo"])

    table_rows = [
        (
            "Прогноз стоимости жилья",
            "Компактная глубокая конфигурация (1 скрытый уровень)",
            "RMSE/R² = "
            + _select_row(reg_rows, "Глубокая конфигурация с одним скрытым уровнем")["rmse"]
            + " / "
            + _select_row(reg_rows, "Глубокая конфигурация с одним скрытым уровнем")["r2"],
            "Градиентный бустинг: RMSE/R² = "
            + _select_row(reg_rows, "Gradient Boosting")["rmse"]
            + " / "
            + _select_row(reg_rows, "Gradient Boosting")["r2"],
        ),
        (
            "Бинарная классификация стоимости жилья",
            "Глубокая конфигурация с сохранением исходных признаков",
            "Accuracy/F1 = "
            + _select_row(cls_rows, "Контекстная глубокая конфигурация")["accuracy"]
            + " / "
            + _select_row(cls_rows, "Контекстная глубокая конфигурация")["f1"],
            "Градиентный бустинг: Accuracy/F1 = "
            + _select_row(cls_rows, "Gradient Boosting")["accuracy"]
            + " / "
            + _select_row(cls_rows, "Gradient Boosting")["f1"],
        ),
        (
            "Бинарная классификация типов покрытия",
            "Компактная глубокая конфигурация (1 скрытый уровень)",
            "Accuracy/F1 = "
            + _select_row(cov_rows, "Глубокая конфигурация с одним скрытым уровнем")["accuracy"]
            + " / "
            + _select_row(cov_rows, "Глубокая конфигурация с одним скрытым уровнем")["f1"],
            "Случайный лес: Accuracy/F1 = "
            + _select_row(cov_rows, "Random Forest")["accuracy"]
            + " / "
            + _select_row(cov_rows, "Random Forest")["f1"],
        ),
    ]

    _add_paragraph(document, "Таблица 1. Ключевые результаты RuFLEX и внешних методов", style="Заголовок табл", align=WD_ALIGN_PARAGRAPH.CENTER)
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    headers = ("Задача", "Лучшая конфигурация RuFLEX", "Результат RuFLEX", "Лучшая внешняя модель")
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(8)

    for row in table_rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cells[i].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            run.font.size = Pt(8)

    document.add_paragraph()


if __name__ == "__main__":
    main()
